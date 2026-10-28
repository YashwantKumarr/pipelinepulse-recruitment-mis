import io
from datetime import UTC, datetime, timedelta

import pytest
from fastapi.testclient import TestClient

from app.config import Settings
from app.main import create_app

FIXED_NOW = datetime(2026, 3, 1, 9, 0, tzinfo=UTC)
WEBHOOK_SECRET = "test-webhook-secret-123"
WEBHOOK_HEADERS = {"X-N8N-Secret": WEBHOOK_SECRET}


class FixedClock:
    """Injected clock so deadline and timestamp behaviour is deterministic."""

    def __init__(self, now: datetime):
        self.now = now

    def __call__(self) -> datetime:
        return self.now

    def advance(self, **kwargs) -> None:
        self.now = self.now + timedelta(**kwargs)


@pytest.fixture
def clock() -> FixedClock:
    return FixedClock(FIXED_NOW)


@pytest.fixture
def settings(tmp_path) -> Settings:
    return Settings(
        database_url=f"sqlite:///{(tmp_path / 'test.db').as_posix()}",
        upload_dir=tmp_path / "uploads",
        enable_n8n_webhooks=True,
        n8n_shared_secret=WEBHOOK_SECRET,
        feedback_sla_hours=72,
    )


@pytest.fixture
def app(settings, clock):
    return create_app(settings=settings, clock=clock)


@pytest.fixture
def client(app):
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture
def make_job(client):
    def _make(**overrides):
        body = {"title": "Data Engineer", "requirements": ["Python", "SQL"], "location": "Remote"}
        body.update(overrides)
        response = client.post("/api/jobs", json=body)
        assert response.status_code == 201, response.text
        return response.json()

    return _make


@pytest.fixture
def make_candidate(client):
    def _make(**overrides):
        body = {"display_name": "Synthetic Person", "source": "referral"}
        body.update(overrides)
        response = client.post("/api/candidates", json=body)
        assert response.status_code == 201, response.text
        return response.json()

    return _make


@pytest.fixture
def upload(client):
    def _upload(filename, data, name="Sample Candidate", content_type="application/octet-stream", **form):
        fields = {"display_name": name, **form}
        return client.post(
            "/api/candidates/upload-resume",
            files={"file": (filename, data, content_type)},
            data=fields,
        )

    return _upload


def make_docx(paragraphs, table_rows=()) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=0, cols=2)
        for left, right in table_rows:
            cells = table.add_row().cells
            cells[0].text = left
            cells[1].text = right
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf(lines) -> bytes:
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    y = 72
    for line in lines:
        page.insert_text((72, y), line)
        y += 18
    data = document.tobytes()
    document.close()
    return data
