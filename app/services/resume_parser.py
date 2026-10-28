"""Resume extraction for PDF, DOCX, TXT and Markdown.

Security notes:
- The client filename is used only to read the extension; it is never used as a path.
- Size is enforced by the caller reading at most max_bytes + 1 bytes, and again here.
- Content is validated by signature (PDF header, ZIP/DOCX structure, UTF-8 text, no NUL bytes).
- Extracted text is untrusted data. It is never executed or interpreted as instructions.
- OCR is NOT implemented: image-only PDFs are rejected with no_extractable_text.
"""

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

import docx
import pymupdf

from app.errors import AppError

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".txt", ".md"})
MAX_TEXT_CHARS = 200_000


class ResumeParseError(AppError):
    pass


@dataclass(frozen=True)
class ParsedResume:
    text: str
    file_format: str


def normalize_text(raw: str) -> str:
    text = unicodedata.normalize("NFC", raw).replace("\r\n", "\n").replace("\r", "\n")
    text = "".join(ch for ch in text if ch in "\n\t" or unicodedata.category(ch)[0] != "C")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_TEXT_CHARS]


def _corrupt() -> ResumeParseError:
    return ResumeParseError(422, "corrupt_document", "The file could not be read as a valid document.")


def _extract_pdf(data: bytes) -> str:
    if not data.startswith(b"%PDF-"):
        raise _corrupt()
    try:
        with pymupdf.open(stream=data, filetype="pdf") as document:
            if document.is_encrypted:
                raise ResumeParseError(
                    422, "encrypted_document", "Password-protected PDFs are not supported."
                )
            return "\n".join(page.get_text("text") for page in document)
    except ResumeParseError:
        raise
    except Exception as exc:  # malformed PDF structures raise library-specific errors
        raise _corrupt() from exc


def _extract_docx(data: bytes) -> str:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise _corrupt()
    try:
        document = docx.Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
    except Exception as exc:
        raise _corrupt() from exc
    return "\n".join(parts)


def _extract_plain(data: bytes) -> str:
    if b"\x00" in data:
        raise ResumeParseError(422, "binary_content", "Text files must not contain binary data.")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ResumeParseError(422, "unsupported_encoding", "Text files must be UTF-8 encoded.") from exc


def parse_resume(filename: str, data: bytes, max_bytes: int) -> ParsedResume:
    extension = PurePath(filename or "").suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeParseError(415, "unsupported_format", "Supported formats: PDF, DOCX, TXT, Markdown.")
    if len(data) > max_bytes:
        raise ResumeParseError(
            413, "file_too_large", f"File exceeds the {max_bytes // (1024 * 1024)} MB limit."
        )
    if not data.strip():
        raise ResumeParseError(422, "empty_file", "The uploaded file is empty.")

    if extension == ".pdf":
        raw = _extract_pdf(data)
    elif extension == ".docx":
        raw = _extract_docx(data)
    else:
        raw = _extract_plain(data)

    text = normalize_text(raw)
    if not text:
        raise ResumeParseError(
            422,
            "no_extractable_text",
            "No text could be extracted. Scanned or image-only files need OCR, which is not supported.",
        )
    return ParsedResume(text=text, file_format=extension.lstrip("."))
